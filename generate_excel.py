import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# Style definitions
font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
font_subtitle = Font(name="Calibri", size=11, italic=True, color="A5B4FC")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
font_regular = Font(name="Calibri", size=10, color="1E293B")
font_price = Font(name="Calibri", size=11, bold=True, color="065F46")
font_badge = Font(name="Calibri", size=10, bold=True, color="4338CA")

fill_title = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid") # Dark Void
fill_header = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid") # Deep Indigo
fill_accent = PatternFill(start_color="06B6D4", end_color="06B6D4", fill_type="solid") # Cyan
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_highlight = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid") # Light Mint

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

header_border = Border(
    left=Side(style='thin', color='312E81'),
    right=Side(style='thin', color='312E81'),
    top=Side(style='medium', color='06B6D4'),
    bottom=Side(style='medium', color='06B6D4')
)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

# -------------------------------------------------------------
# SHEET 1: Core Packages
# -------------------------------------------------------------
ws1 = wb.active
ws1.title = "Core Packages"
ws1.views.sheetView[0].showGridLines = True

# Title Block
ws1.merge_cells("A1:F1")
ws1["A1"] = "RYNEX — WEB & SOFTWARE SOLUTIONS"
ws1["A1"].font = font_title
ws1["A1"].fill = fill_title
ws1["A1"].alignment = align_center

ws1.merge_cells("A2:F2")
ws1["A2"] = "Official Cost Structure Rate Card • Starting ₹3,000 to ₹19,999 Maximum Cap • Chennai Hub"
ws1["A2"].font = font_subtitle
ws1["A2"].fill = fill_title
ws1["A2"].alignment = align_center

# Headers
headers_1 = ["Package Tier", "Base Price (INR)", "Scope / Pages", "Target Client", "Key Deliverables", "Turnaround"]
for col_num, h in enumerate(headers_1, 1):
    cell = ws1.cell(row=4, column=col_num)
    cell.value = h
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = header_border

rows_1 = [
    ("Starter Modern Website", 3000, "1 Single Page (Landing/Portfolio)", "Freelancers, Creators, Portfolios, Quick Launch", "Responsive Layout, Hero CTA, Lead Form, WhatsApp Chat, Fast Load (<1s)", "3 - 5 Days"),
    ("Multi-Page Business Website", 6000, "3 - 5 Full Pages", "Small Businesses, Agencies, Consultancies, Clinics", "Home, About, Services, Contact, Full SEO, Custom UI/UX, Dynamic Forms", "7 - 10 Days"),
    ("Dynamic Web Application", 10000, "5 - 8 Views / Dashboards", "SaaS Startups, Client Portals, Booking Tools", "React/Next.js UI, Node/Python API, Database, Filter/Search, Client Portal", "12 - 16 Days"),
    ("Full E-Commerce / Custom SaaS", 15000, "Complete Multi-Module System", "Online Stores, Digital Products, Enterprise SaaS", "Product Catalog, Cart, Order Management, Payments, Admin Panel, Edge CDN (Max ₹19,999)", "2 - 3 Weeks")
]

for row_idx, r in enumerate(rows_1, 5):
    ws1.cell(row=row_idx, column=1, value=r[0]).font = font_bold
    ws1.cell(row=row_idx, column=2, value=r[1]).font = font_price
    ws1.cell(row=row_idx, column=2).number_format = '₹#,##0'
    ws1.cell(row=row_idx, column=2).alignment = align_right
    ws1.cell(row=row_idx, column=3, value=r[2]).font = font_regular
    ws1.cell(row=row_idx, column=4, value=r[3]).font = font_regular
    ws1.cell(row=row_idx, column=5, value=r[4]).font = font_regular
    ws1.cell(row=row_idx, column=6, value=r[5]).font = font_bold
    ws1.cell(row=row_idx, column=6).alignment = align_center

    for c in range(1, 7):
        cell = ws1.cell(row=row_idx, column=c)
        cell.border = thin_border
        if row_idx % 2 == 0:
            cell.fill = fill_zebra

# -------------------------------------------------------------
# SHEET 2: Modular Add-ons
# -------------------------------------------------------------
ws2 = wb.create_sheet(title="Modular Add-ons")
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells("A1:E1")
ws2["A1"] = "RYNEX — MODULAR CAPABILITY ADD-ONS"
ws2["A1"].font = font_title
ws2["A1"].fill = fill_title
ws2["A1"].alignment = align_center

ws2.merge_cells("A2:E2")
ws2["A2"] = "Add standalone high-performance features to any package (All-inclusive max rate ₹19,999)"
ws2["A2"].font = font_subtitle
ws2["A2"].fill = fill_title
ws2["A2"].alignment = align_center

headers_2 = ["Add-on Name", "Additional Cost (INR)", "Feature Category", "Technical Description & Deliverables", "Delivery Impact"]
for col_num, h in enumerate(headers_2, 1):
    cell = ws2.cell(row=4, column=col_num)
    cell.value = h
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = header_border

rows_2 = [
    ("User Auth & Profiles", 1000, "Security & Accounts", "Sign up / Login, Password recovery, JWT Sessions, User Profiles", "+1 - 2 Days"),
    ("Payment Gateway Integration", 1000, "Fintech & Checkout", "UPI, QR Code, Credit/Debit cards, Net Banking (Razorpay / Stripe)", "+1 - 2 Days"),
    ("Admin Control & Analytics Dashboard", 1000, "Management", "Admin portal with charts, visitor counts, enquiry logs & data export", "+2 - 3 Days"),
    ("3D Interactive FX & WebGL", 1000, "Visuals & UI", "Custom 3D canvas animations, gyro tilt, mouse particle shaders", "+1 - 2 Days"),
    ("AI Chatbot / Smart Assistant", 1500, "Artificial Intelligence", "Custom AI chatbot trained on business FAQs (OpenAI / Gemini)", "+2 - 3 Days"),
    ("Custom Domain Setup & Fast Edge CDN", 500, "DevOps & Cloud", "Custom domain linking, SSL certificate (HTTPS), Cloudflare CDN boost", "+1 Day")
]

for row_idx, r in enumerate(rows_2, 5):
    ws2.cell(row=row_idx, column=1, value=r[0]).font = font_bold
    ws2.cell(row=row_idx, column=2, value=r[1]).font = font_price
    ws2.cell(row=row_idx, column=2).number_format = '₹#,##0'
    ws2.cell(row=row_idx, column=2).alignment = align_right
    ws2.cell(row=row_idx, column=3, value=r[2]).font = font_badge
    ws2.cell(row=row_idx, column=4, value=r[3]).font = font_regular
    ws2.cell(row=row_idx, column=5, value=r[4]).font = font_regular
    ws2.cell(row=row_idx, column=5).alignment = align_center

    for c in range(1, 6):
        cell = ws2.cell(row=row_idx, column=c)
        cell.border = thin_border
        if row_idx % 2 == 0:
            cell.fill = fill_zebra

# -------------------------------------------------------------
# SHEET 3: Sample Bundles & Estimator
# -------------------------------------------------------------
ws3 = wb.create_sheet(title="Sample Bundles")
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells("A1:E1")
ws3["A1"] = "RYNEX — POPULAR PROJECT BUNDLES"
ws3["A1"].font = font_title
ws3["A1"].fill = fill_title
ws3["A1"].alignment = align_center

ws3.merge_cells("A2:E2")
ws3["A2"] = "Curated ready-to-launch bundles with guaranteed maximum pricing cap of ₹19,999"
ws3["A2"].font = font_subtitle
ws3["A2"].fill = fill_title
ws3["A2"].alignment = align_center

headers_3 = ["Bundle Name", "Base Package", "Included Add-ons", "Calculated Value", "Final Capped Rate (INR)"]
for col_num, h in enumerate(headers_3, 1):
    cell = ws3.cell(row=4, column=col_num)
    cell.value = h
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = header_border

rows_3 = [
    ("Budget Kickstart", "Starter Modern Website (₹3,000)", "Core Landing Page & WhatsApp Lead Gen", 3000, 3000),
    ("Professional Business", "Multi-Page Business Site (₹6,000)", "Domain Setup & Fast CDN (+₹500)", 6500, 6500),
    ("Interactive Brand Agency", "Multi-Page Business Site (₹6,000)", "3D Interactive FX (+₹1,000) + AI Chatbot (+₹1,500)", 8500, 8500),
    ("Modern E-Commerce Store", "Full E-Commerce Platform (₹15,000)", "Payment Gateway (+₹1,000) + User Auth (+₹1,000)", 17000, 17000),
    ("Ultimate Full-Stack SaaS", "Full SaaS Platform (₹15,000)", "All 6 Add-ons (Auth, Payments, Admin, 3D, AI, CDN)", 21000, 19999)
]

for row_idx, r in enumerate(rows_3, 5):
    ws3.cell(row=row_idx, column=1, value=r[0]).font = font_bold
    ws3.cell(row=row_idx, column=2, value=r[1]).font = font_regular
    ws3.cell(row=row_idx, column=3, value=r[2]).font = font_regular
    ws3.cell(row=row_idx, column=4, value=r[3]).font = font_regular
    ws3.cell(row=row_idx, column=4).number_format = '₹#,##0'
    ws3.cell(row=row_idx, column=4).alignment = align_right
    
    cell_final = ws3.cell(row=row_idx, column=5, value=r[4])
    cell_final.font = font_price
    cell_final.number_format = '₹#,##0'
    cell_final.alignment = align_right
    cell_final.fill = fill_highlight

    for c in range(1, 6):
        cell = ws3.cell(row=row_idx, column=c)
        cell.border = thin_border

# Auto-adjust column widths for all sheets
for ws in [ws1, ws2, ws3]:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row in [1, 2]:
                continue
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

ws1.row_dimensions[1].height = 26
ws1.row_dimensions[2].height = 20
ws1.row_dimensions[4].height = 24
for r in range(5, 9):
    ws1.row_dimensions[r].height = 28

ws2.row_dimensions[1].height = 26
ws2.row_dimensions[2].height = 20
ws2.row_dimensions[4].height = 24
for r in range(5, 11):
    ws2.row_dimensions[r].height = 26

ws3.row_dimensions[1].height = 26
ws3.row_dimensions[2].height = 20
ws3.row_dimensions[4].height = 24
for r in range(5, 10):
    ws3.row_dimensions[r].height = 26

wb.save("c:/Users/Barath/OneDrive/Documents/Freelanzing/rynex/RYNEX_Pricing_Structure.xlsx")
print("Excel spreadsheet generated successfully!")
